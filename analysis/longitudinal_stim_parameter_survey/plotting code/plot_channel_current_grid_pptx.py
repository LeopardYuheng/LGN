"""
plot_channel_current_grid_pptx.py
----------------------------------
Compares one or more channels' dF/F response, at one user-chosen time
point, across all of each channel's stimulation currents from a step-5_A
output folder — and assembles the result into a single PPTX with ONE
SLIDE PER SELECTED CHANNEL.

Layout: up to 6 panels (2 rows x 3 columns) per slide, one per current
(ascending), all sharing ONE user-defined colorbar range (vmin/vmax, same
range across every channel/slide) so the panels are directly comparable.

Optional overlay: the user can choose to overlay brain-area boundaries
(V1 + all higher visual areas) from a day_setup.mat produced by the
UPDATED step 6 (retino_alignment_with_brain_mask_6.m). Those boundary
fields (V1_mask_stim, retOverlay_stim) already share the same pixel grid
as the step-5_A dF/F movies, so no extra warping happens here -- V1 is
drawn as a gold contour outline, all other area boundaries as an orange
overlay, on every panel of every slide.

Orientation note:
    mean_dff_movie is read back from the MATLAB v7.3 (HDF5) .mat file with
    numpy/h5py, whose dimension order is the reverse of MATLAB's. It is
    transposed back to MATLAB's (H, W, T) indexing, and each frame is drawn
    with imshow(..., origin='lower') — this is the matplotlib equivalent of
    MATLAB's set(ax,'YDir','normal'), which is what every dF/F figure in
    this pipeline (step 5_A, 7_A, ...) uses. So this figure's orientation
    matches the step-5_A output figures exactly.

Usage:
    pip install python-pptx pillow matplotlib h5py numpy scipy
    python plot_channel_current_grid_pptx.py
"""

import os
import re
import tkinter as tk
from tkinter import filedialog, simpledialog, messagebox

import numpy as np
import h5py
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from scipy.ndimage import binary_dilation

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN


_CH_CUR_RE = re.compile(r'mean_dff_ch(\d+)_([\d.]+)uA\.mat$', re.I)

_SLIDE_W = 13.333
_SLIDE_H = 7.5
_N_COLS  = 3
_N_ROWS  = 2


def load_mean_dff(mat_path):
    """Load mean_dff_movie (H,W,T), t_s (T,), final_mask (H,W) or None from
    a step-5_A mean_dff_ch{N}_{I}uA.mat file (MATLAB v7.3 / HDF5)."""
    with h5py.File(mat_path, 'r') as f:
        raw = np.array(f['mean_dff_movie'])           # stored as (T, W, H)
        movie = np.transpose(raw, (2, 1, 0))           # -> MATLAB's (H, W, T)
        t_s = np.array(f['t_s']).ravel().astype(float)
        final_mask = None
        if 'final_mask' in f:
            final_mask = np.array(f['final_mask']).T.astype(bool)  # -> (H, W)
    return movie, t_s, final_mask


def find_channel_current_files(root):
    """Recursively scan root for mean_dff_ch{N}_{I}uA.mat files.
    Returns dict: channel(int) -> list of (current(float), path), sorted by current."""
    by_channel = {}
    for dirpath, _, filenames in os.walk(root):
        for fn in filenames:
            m = _CH_CUR_RE.search(fn)
            if not m:
                continue
            ch = int(m.group(1))
            cur = float(m.group(2))
            by_channel.setdefault(ch, []).append((cur, os.path.join(dirpath, fn)))
    for ch in by_channel:
        by_channel[ch].sort(key=lambda x: x[0])
    return by_channel


def nearest_frame_idx(t_s, t_target):
    return int(np.argmin(np.abs(t_s - t_target)))


def load_day_setup_boundaries(ds_path):
    """Load V1_mask_stim and retOverlay_stim (all-area boundaries) from a
    day_setup.mat (MATLAB v7.3 / HDF5, from the updated step 6). Both are
    already on the same pixel grid as the step-5_A dF/F movies. Returns
    (V1_mask, region_bound) as (H,W) bool arrays -- either may be None if
    missing or unreadable."""
    V1_mask = None
    region_bound = None
    try:
        with h5py.File(ds_path, 'r') as f:
            ra = f['day_setup']['retino_align']
            if 'V1_mask_stim' in ra:
                V1_mask = np.array(ra['V1_mask_stim']).T.astype(bool)   # -> (H, W)
            if 'retOverlay_stim' in ra:
                region_bound = np.array(ra['retOverlay_stim']).T.astype(bool)  # -> (H, W)
    except Exception as e:
        print(f'Warning: could not read boundaries from day_setup.mat ({ds_path}): {e}')
    if V1_mask is None:
        print('Warning: day_setup has no V1_mask_stim -- V1 outline will not be shown.')
    if region_bound is None:
        print('Warning: day_setup has no retOverlay_stim -- visual-area boundaries will not be shown.')
    return V1_mask, region_bound


def boundary_rgba_overlay(mask, color_rgb):
    """RGBA image: fully opaque `color_rgb` where mask is True, fully
    transparent elsewhere -- for drawing a thin boundary mask on top of an
    imshow'd frame without altering its own colormap/values."""
    h, w = mask.shape
    rgba = np.zeros((h, w, 4), dtype=float)
    rgba[..., 0] = color_rgb[0]
    rgba[..., 1] = color_rgb[1]
    rgba[..., 2] = color_rgb[2]
    rgba[..., 3] = mask.astype(float)
    return rgba


def main():
    root_tk = tk.Tk()
    root_tk.withdraw()

    root_dir = filedialog.askdirectory(
        title='Select step-5_A output folder (contains ch{N}_{I}uA/ subfolders)')
    if not root_dir:
        print('No folder selected. Exiting.')
        return

    by_channel = find_channel_current_files(root_dir)
    if not by_channel:
        messagebox.showerror('No files found',
                              f'No mean_dff_ch*.mat files found under:\n{root_dir}')
        return

    channels = sorted(by_channel.keys())
    ch_str = ', '.join(str(c) for c in channels)
    ch_answer = simpledialog.askstring(
        'Select channel(s)',
        f'Available channels: {ch_str}\n\n'
        'Enter channel number(s) to analyze, comma/space-separated (e.g. 30, 35, 43):')
    if not ch_answer:
        print('No channel(s) specified. Exiting.')
        return

    sel_channels = []
    for tok in re.split(r'[,\s]+', ch_answer.strip()):
        if not tok:
            continue
        try:
            ch = int(tok)
        except ValueError:
            print(f'Ignoring invalid channel token: "{tok}"')
            continue
        if ch not in by_channel:
            print(f'Ignoring unknown channel: {ch}')
            continue
        if ch not in sel_channels:
            sel_channels.append(ch)

    if not sel_channels:
        messagebox.showerror('No valid channels', 'No valid channel numbers were selected.')
        return
    print(f'Selected {len(sel_channels)} channel(s): '
          + ', '.join(str(c) for c in sel_channels))

    t_answer = simpledialog.askfloat(
        'Time point', 'Time point t (s) to display, e.g. 0.5:', initialvalue=0.5)
    if t_answer is None:
        print('No time point specified. Exiting.')
        return

    vmin = simpledialog.askfloat('Color scale', 'Colorbar LOWER limit (dF/F):', initialvalue=-0.02)
    if vmin is None:
        print('No lower limit specified. Exiting.')
        return
    vmax = simpledialog.askfloat('Color scale', 'Colorbar UPPER limit (dF/F):', initialvalue=0.02)
    if vmax is None:
        print('No upper limit specified. Exiting.')
        return
    if not (vmax > vmin):
        messagebox.showerror('Invalid range', 'Upper limit must be greater than lower limit.')
        return

    V1_mask = None
    region_bound = None
    overlay_answer = messagebox.askyesno(
        'Visual area boundaries',
        'Overlay brain area boundaries (V1 + higher visual areas) from a day_setup.mat?')
    if overlay_answer:
        ds_path = filedialog.askopenfilename(
            title='Select day_setup.mat (from step 6) -- provides V1 + visual-area boundaries',
            filetypes=[('MAT files', '*.mat'), ('All files', '*.*')])
        if ds_path:
            V1_mask, region_bound = load_day_setup_boundaries(ds_path)
        else:
            print('No day_setup.mat selected -- continuing without boundary overlay.')

    save_dir = filedialog.askdirectory(
        title='Select output folder for the PNG/PPTX', initialdir=root_dir)
    if not save_dir:
        print('No output folder selected. Exiting.')
        return

    cmap = matplotlib.colormaps['jet'].copy()
    cmap.set_bad(color=[0.15, 0.15, 0.15])   # outside-brain / masked-out pixels

    prs = Presentation()
    prs.slide_width = Inches(_SLIDE_W)
    prs.slide_height = Inches(_SLIDE_H)

    png_paths = []
    for ch in sel_channels:
        entries = by_channel[ch]   # list of (current, path), ascending
        n_cur = len(entries)
        print(f'\nChannel {ch}: {n_cur} current condition(s) found: '
              + ', '.join(f'{c:g} uA' for c, _ in entries))

        max_panels = _N_ROWS * _N_COLS
        if n_cur > max_panels:
            print(f'Note: {n_cur} currents found but only the first {max_panels} '
                  f'(lowest currents) will be shown in the {_N_ROWS}x{_N_COLS} grid.')
            entries = entries[:max_panels]
            n_cur = len(entries)

        png_path, t_label = render_channel_grid_png(ch, entries, n_cur, t_answer, vmin, vmax, cmap, save_dir,
                                                     V1_mask, region_bound)
        png_paths.append(png_path)
        add_channel_slide(prs, ch, png_path, t_label, vmin, vmax)

    ch_tag = '-'.join(str(c) for c in sel_channels)
    pptx_fname = f'ch{ch_tag}_current_grid_t{t_answer:g}s.pptx'
    pptx_path = os.path.join(save_dir, pptx_fname)
    prs.save(pptx_path)
    print(f'\nSaved PPTX ({len(sel_channels)} slide(s)): {pptx_path}')
    messagebox.showinfo('Done', f'Saved {len(sel_channels)} slide(s) to:\n{pptx_path}')


def render_channel_grid_png(ch, entries, n_cur, t_answer, vmin, vmax, cmap, save_dir,
                             V1_mask=None, region_bound=None):
    """Renders one channel's current-comparison grid figure and saves it as
    a PNG. Returns (png_path, t_label) where t_label is the actual time (s)
    of the nearest frame used. If V1_mask/region_bound are given (from a
    day_setup.mat), they are overlaid on every panel: V1 as a gold contour
    outline, all other visual-area boundaries as a dilated orange overlay."""
    region_bound_disp = None
    if region_bound is not None:
        region_bound_disp = binary_dilation(region_bound, iterations=1)

    fig, axes = plt.subplots(_N_ROWS, _N_COLS, figsize=(13, 8), constrained_layout=True)
    axes = np.atleast_2d(axes)
    im_last = None
    actual_t_used = []

    for idx in range(_N_ROWS * _N_COLS):
        r, c = divmod(idx, _N_COLS)
        ax = axes[r, c]
        if idx >= n_cur:
            ax.axis('off')
            continue

        cur, mat_path = entries[idx]
        movie, t_s, final_mask = load_mean_dff(mat_path)
        fi = nearest_frame_idx(t_s, t_answer)
        frame = movie[:, :, fi].astype(float)
        actual_t_used.append(t_s[fi])

        if final_mask is None:
            final_mask = ~np.all(np.isnan(movie), axis=2)
        frame = np.where(final_mask, frame, np.nan)

        im = ax.imshow(frame, cmap=cmap, vmin=vmin, vmax=vmax, origin='lower')

        if region_bound_disp is not None:
            if region_bound_disp.shape == frame.shape:
                ax.imshow(boundary_rgba_overlay(region_bound_disp, (1.0, 0.55, 0.0)), origin='lower')
            else:
                print(f'Warning: retOverlay_stim shape {region_bound_disp.shape} does not match '
                      f'frame shape {frame.shape} -- skipping boundary overlay for {mat_path}')
        if V1_mask is not None:
            if V1_mask.shape == frame.shape:
                ax.contour(V1_mask.astype(float), levels=[0.5], colors=['gold'], linewidths=1.3, origin='lower')
            else:
                print(f'Warning: V1_mask_stim shape {V1_mask.shape} does not match '
                      f'frame shape {frame.shape} -- skipping V1 outline for {mat_path}')

        ax.set_xticks([]); ax.set_yticks([])
        ax.set_title(f'{cur:g} μA', fontsize=13, fontweight='bold')
        im_last = im

    if im_last is not None:
        cb = fig.colorbar(im_last, ax=axes, shrink=0.85, pad=0.02)
        cb.set_label('ΔF/F', fontsize=12)
        cb.ax.tick_params(labelsize=10)

    t_label = actual_t_used[0] if actual_t_used else t_answer
    fig.suptitle(f'Channel {ch}  |  t = {t_label:+.2f} s  |  '
                 f'color scale [{vmin:g}, {vmax:g}]', fontsize=15, fontweight='bold')

    png_fname = f'ch{ch}_current_grid_t{t_answer:g}s.png'
    png_path = os.path.join(save_dir, png_fname)
    fig.savefig(png_path, dpi=150)
    plt.close(fig)
    print(f'Saved figure: {png_path}')
    return png_path, t_label


def add_channel_slide(prs, ch, png_path, t_label, vmin, vmax):
    """Adds one slide to prs showing png_path (one channel's current-grid
    figure) with a title textbox."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    tb = slide.shapes.add_textbox(Inches(0.3), Inches(0.1), Inches(_SLIDE_W - 0.6), Inches(0.5))
    para = tb.text_frame.paragraphs[0]
    para.alignment = PP_ALIGN.LEFT
    run = para.add_run()
    run.text = (f'Channel {ch}  |  dF/F at t = {t_label:+.2f} s  |  '
                f'color scale [{vmin:g}, {vmax:g}]')
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x20, 0x20, 0x20)

    from PIL import Image as PILImage
    w_px, h_px = PILImage.open(png_path).size
    ar = w_px / h_px
    img_w = Inches(_SLIDE_W - 0.6)
    img_h = img_w / ar
    max_h = Inches(_SLIDE_H - 0.9)
    if img_h > max_h:
        img_h = max_h
        img_w = img_h * ar
    img_x = (Inches(_SLIDE_W) - img_w) // 2
    img_y = Inches(0.7)
    slide.shapes.add_picture(png_path, img_x, img_y, width=img_w, height=img_h)


if __name__ == '__main__':
    main()