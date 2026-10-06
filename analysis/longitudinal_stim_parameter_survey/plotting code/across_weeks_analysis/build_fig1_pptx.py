"""
build_fig1_pptx.py
------------------
FIGURE 1 deck — one slide per effective channel.

Reads the PNGs written by make_fig1_frame_grids.m and assembles a 16:9
PowerPoint with one slide per (mouse, channel), each showing that channel's
frame grid across weeks.

Each MATLAB figure already contains every week as a row, so the normal case
is one image per slide, scaled to fill. If a channel happens to have several
PNGs (e.g. you rendered one per week instead), they are stacked vertically
on the same slide in week order, so this works either way.

Usage:
    pip install python-pptx pillow
    python build_fig1_pptx.py [figure1_frame_grids folder]

With no argument it opens a folder picker.

Output:
    <folder>/figure1_frame_grids.pptx
"""

import os
import re
import sys
from collections import defaultdict

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# --- slide geometry (inches), matching build_avg_results_pptx.py ----------
SLIDE_W = 13.333
SLIDE_H = 7.5
MARGIN = 0.25
TITLE_H = 0.62
V_GAP = 0.06

# fig1_frame_grid_<MOUSE>_ch<N>_<I>uA.png
FNAME_RE = re.compile(r'fig1_frame_grid_(?P<mouse>\w+?)_ch(?P<ch>\d+)_(?P<cur>[\d.]+)uA'
                      r'(?:_(?P<week>[\w+-]+))?\.png$', re.I)

# Sort key for the optional per-week variant: pre W4 before post W1.
WEEK_RE = re.compile(r'(pre|post)\s*W?(\d+)', re.I)


def week_sort_key(week_tag):
    if not week_tag:
        return 0
    m = WEEK_RE.search(week_tag)
    if not m:
        return 0
    phase, k = m.group(1).lower(), int(m.group(2))
    return k - 5 if phase == 'pre' else k


def collect(folder):
    """Group PNGs by (mouse, channel, current) -> [(week_key, path), ...]."""
    groups = defaultdict(list)
    for fn in sorted(os.listdir(folder)):
        m = FNAME_RE.match(fn)
        if not m:
            continue
        key = (m.group('mouse'), int(m.group('ch')), float(m.group('cur')))
        groups[key].append((week_sort_key(m.group('week')), os.path.join(folder, fn)))
    for key in groups:
        groups[key].sort(key=lambda t: t[0])
    return groups


def image_ar(png):
    """width/height of a PNG, falling back to a wide default."""
    try:
        from PIL import Image
        w, h = Image.open(png).size
        return w / h
    except Exception:
        return 16 / 6


def add_slide(prs, title_text, images):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    tb = slide.shapes.add_textbox(
        Inches(MARGIN), Inches(0.08),
        Inches(SLIDE_W - 2 * MARGIN), Inches(TITLE_H - 0.08),
    )
    para = tb.text_frame.paragraphs[0]
    para.alignment = PP_ALIGN.LEFT
    run = para.add_run()
    run.text = title_text
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x20, 0x20, 0x20)

    # Area available to the image stack
    area_w = SLIDE_W - 2 * MARGIN
    area_h = SLIDE_H - TITLE_H - MARGIN
    n = len(images)
    slot_h = (area_h - (n - 1) * V_GAP) / n

    y = TITLE_H
    for png in images:
        ar = image_ar(png)
        w = area_w
        h = w / ar
        if h > slot_h:            # too tall for its slot — fit by height
            h = slot_h
            w = h * ar
        x = MARGIN + (area_w - w) / 2
        slide.shapes.add_picture(png, Inches(x), Inches(y + (slot_h - h) / 2),
                                 width=Inches(w), height=Inches(h))
        y += slot_h + V_GAP


def main():
    if len(sys.argv) > 1:
        folder = sys.argv[1]
    else:
        import tkinter as tk
        from tkinter import filedialog
        tk.Tk().withdraw()
        folder = filedialog.askdirectory(title='Select the figure1_frame_grids folder')
    if not folder:
        print('No folder selected. Exiting.')
        return
    if not os.path.isdir(folder):
        print(f'Not a folder: {folder}')
        return

    groups = collect(folder)
    if not groups:
        print(f'No fig1_frame_grid_*.png found in:\n  {folder}\n'
              'Run make_fig1_frame_grids.m first.')
        return

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    for n_slide, (mouse, ch, cur) in enumerate(
            sorted(groups, key=lambda k: (k[0], k[1])), start=1):
        images = [p for _, p in groups[(mouse, ch, cur)]]
        title = f'{mouse}  ch{ch}  |  {cur:g} µA  |  frame grid across weeks'
        add_slide(prs, title, images)
        print(f'  [slide {n_slide:2d}] {title}  ({len(images)} image)')

    out_path = os.path.join(folder, 'figure1_frame_grids.pptx')
    prs.save(out_path)
    print(f'\nSaved {len(groups)} slides -> {out_path}')


if __name__ == '__main__':
    main()
