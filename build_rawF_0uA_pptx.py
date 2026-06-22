"""
build_rawF_0uA_pptx.py
----------------------
Assembles the raw-F 0 uA sanity-check frame-grids produced by
plot_raw_F_0uA_from_tiff.m into a single PowerPoint deck (one slide per trial,
preceded by the per-channel AVERAGE slide), like the 9_C / 9_D decks.

1. Run plot_raw_F_0uA_from_tiff.m in MATLAB and note its output folder
   (it contains ch{N}_0uA/ subfolders of PNGs).
2. Set ROOT below to that folder.
3. pip install python-pptx   (once)
4. python build_rawF_0uA_pptx.py
"""
import os, re, glob
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# Folder that plot_raw_F_0uA_from_tiff.m wrote into (contains ch*_0uA/ subfolders).
ROOT = r"C:\Project\LGN_analysis_results\LGN11_20260326_experiment\rawF_0uA_sanitycheck"
PPTX = os.path.join(ROOT, "rawF_0uA_sanitycheck.pptx")

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def cond_of(d):
    """Return (channel, current) parsed from a 'ch{N}_{I}uA' folder name."""
    m = re.search(r"ch(\d+)_([\d.]+)uA", os.path.basename(d))
    return (int(m.group(1)), float(m.group(2))) if m else (1_000_000, 0.0)


def trial_of(png):
    m = re.search(r"trial(\d+)", os.path.basename(png))
    return int(m.group(1)) if m else -1


def add_slide(title, png):
    slide = prs.slides.add_slide(blank)
    tb = slide.shapes.add_textbox(Inches(0.2), Inches(0.1), prs.slide_width - Inches(0.4), Inches(0.6))
    p = tb.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    r = p.add_run(); r.text = title
    r.font.size = Pt(20); r.font.bold = True; r.font.color.rgb = RGBColor(0x20, 0x20, 0x20)
    # fit the frame-grid under the title
    avail_w = prs.slide_width - Inches(0.4)
    avail_h = prs.slide_height - Inches(0.8)
    try:
        from PIL import Image
        iw, ih = Image.open(png).size
        ar = iw / ih
    except Exception:
        ar = 1150 / 1000   # fallback to the frame-grid's nominal aspect
    w = avail_w; h = w / ar
    if h > avail_h:
        h = avail_h; w = h * ar
    x = (prs.slide_width - w) / 2
    y = Inches(0.75) + (avail_h - h) / 2
    slide.shapes.add_picture(png, x, y, width=w, height=h)


cond_dirs = sorted([d for d in glob.glob(os.path.join(ROOT, "ch*uA")) if os.path.isdir(d)],
                   key=cond_of)
if not cond_dirs:
    raise SystemExit(f"No ch*_*uA subfolders found in {ROOT}")

n = 0
for d in cond_dirs:
    ch, cur = cond_of(d)
    cur_lbl = f"{cur:g}"
    avg = os.path.join(d, "rawF_ch%d_%suA_AVERAGE.png" % (ch, cur_lbl))
    if os.path.isfile(avg):
        add_slide("Channel %d - %s uA - AVERAGE (raw F)" % (ch, cur_lbl), avg)
        n += 1
    pattern = os.path.join(d, "rawF_ch%d_%suA_trial*.png" % (ch, cur_lbl))
    for png in sorted(glob.glob(pattern), key=trial_of):
        add_slide("Channel %d - %s uA - trial %d (raw F)" % (ch, cur_lbl, trial_of(png)), png)
        n += 1

prs.save(PPTX)
print(f"Saved {PPTX} with {n} slides")
